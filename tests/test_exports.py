import asyncio
import io
import json
import unittest
from datetime import date
from unittest.mock import patch
from urllib.parse import unquote

from fastapi import HTTPException
from openpyxl import load_workbook

from backend import main


async def stream_bytes(response):
    return b"".join([chunk async for chunk in response.body_iterator])


async def request_export(path, query):
    messages = []
    scope = {
        "type": "http", "asgi": {"version": "3.0", "spec_version": "2.4"},
        "http_version": "1.1", "method": "GET", "scheme": "http",
        "path": path, "raw_path": path.encode(), "query_string": query.encode(),
        "headers": [], "client": ("test", 50000), "server": ("test", 80),
        "root_path": "",
    }

    async def receive():
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(message):
        messages.append(message)

    await main.app(scope, receive, send)
    status = next(m["status"] for m in messages if m["type"] == "http.response.start")
    body = b"".join(m.get("body", b"") for m in messages if m["type"] == "http.response.body")
    return status, body


def reservation(day, state="예약완료", start="09:00"):
    return {
        "rrState": state, "rrStartTime": start,
        "rpName": "MOOC 스튜디오", "rrBooker": day, "rrPurpose": "테스트",
    }


class ExportPeriodTests(unittest.TestCase):
    def test_month_includes_leap_day(self):
        self.assertEqual(
            main._resolve_export_period(2024, 2, None, None),
            (date(2024, 2, 1), date(2024, 2, 29), False),
        )

    def test_default_is_current_month(self):
        start, end, is_range = main._resolve_export_period(None, None, None, None)
        today = date.today()
        self.assertEqual((start.year, start.month, start.day), (today.year, today.month, 1))
        self.assertEqual(end.month, today.month)
        self.assertFalse(is_range)

    def test_same_day_range_allowed(self):
        day = date(2026, 9, 30)
        self.assertEqual(main._resolve_export_period(None, None, day, day), (day, day, True))

    def test_invalid_periods_rejected(self):
        cases = [
            (None, None, date(2026, 10, 1), date(2026, 9, 30)),
            (None, None, date(2026, 9, 30), None),
            (None, None, None, date(2026, 9, 30)),
            (2026, 0, None, None), (2026, 13, None, None), (0, 1, None, None),
        ]
        for args in cases:
            with self.subTest(args=args), self.assertRaises(HTTPException) as error:
                main._resolve_export_period(*args)
            self.assertEqual(error.exception.status_code, 400)

    def test_fetch_includes_both_ends_across_years(self):
        with patch.object(main, "_fetch", side_effect=lambda day: [reservation(day)]) as fetch:
            data = main._fetch_period(date(2025, 12, 31), date(2026, 1, 2))
        self.assertEqual(sorted(data), [date(2025, 12, 31), date(2026, 1, 1), date(2026, 1, 2)])
        self.assertEqual(sorted(call.args[0] for call in fetch.call_args_list),
                         ["2025-12-31", "2026-01-01", "2026-01-02"])

    def test_fetch_filters_and_sorts_reservations(self):
        day = date(2026, 9, 30)
        items = [reservation("late", start="15:00"), reservation("cancelled", "취소"),
                 reservation("early", start="08:00")]
        with patch.object(main, "_fetch", return_value=items):
            data = main._fetch_period(day, day)
        self.assertEqual([item["rrBooker"] for item in data[day]], ["early", "late"])

    def test_excel_and_txt_keep_same_day_number_in_different_months(self):
        start, end = date(2026, 8, 1), date(2026, 9, 1)

        def fetch(day):
            return [reservation(day)] if day in {"2026-08-01", "2026-09-01"} else []

        with patch.object(main, "_fetch", side_effect=fetch):
            excel = main.export_excel(start_date=start, end_date=end)
            txt = main.export_txt(start_date=start, end_date=end)

        workbook = load_workbook(io.BytesIO(asyncio.run(stream_bytes(excel))))
        values = [str(cell.value) for row in workbook.active for cell in row if cell.value is not None]
        self.assertIn("2026-08-01", values)
        self.assertIn("2026-09-01", values)
        self.assertTrue(any("총 2건" in value for value in values))
        self.assertTrue(any("2026.08.01 ~ 2026.09.01" in value for value in values))
        text = asyncio.run(stream_bytes(txt)).decode("utf-8-sig")
        self.assertIn("2026-08-01", text)
        self.assertIn("2026-09-01", text)
        self.assertIn("2026.08.01 ~ 2026.09.01", text)
        self.assertIn("2026-08-01_2026-09-01.xlsx", unquote(excel.headers["content-disposition"]))
        self.assertIn("2026-08-01_2026-09-01.txt", unquote(txt.headers["content-disposition"]))

    def test_monthly_filename_and_empty_export_preserved(self):
        with patch.object(main, "_fetch", return_value=[]):
            excel = main.export_excel(year=2024, month=2)
            txt = main.export_txt(year=2024, month=2)
        workbook = load_workbook(io.BytesIO(asyncio.run(stream_bytes(excel))))
        self.assertEqual(workbook.active.title, "2024년 2월")
        self.assertEqual(workbook.active["A1"].value, "2024년 2월 예약 내역 없음")
        self.assertIn("2024년02월.xlsx", unquote(excel.headers["content-disposition"]))
        self.assertIn("예약 내역 없음", asyncio.run(stream_bytes(txt)).decode("utf-8-sig"))

    def test_http_rejects_bad_dates_before_fetch(self):
        queries = [
            ("start_date=2026-02-30&end_date=2026-03-01", 422),
            ("start_date=2026-10-01&end_date=2026-09-30", 400),
            ("start_date=2026-09-30", 400),
            ("year=2026&month=13", 400),
        ]
        with patch.object(main, "_fetch") as fetch:
            for path in ("/export", "/export/txt"):
                for query, expected in queries:
                    with self.subTest(path=path, query=query):
                        status, body = asyncio.run(request_export(path, query))
                        self.assertEqual(status, expected)
                        self.assertIn("detail", json.loads(body))
            fetch.assert_not_called()

    def test_http_downloads_use_range_or_month_parameters(self):
        cases = [
            ("start_date=2026-09-30&end_date=2026-09-30", 1, "2026-09-30"),
            ("year=2024&month=2", 29, "2024-02-29"),
        ]
        for path in ("/export", "/export/txt"):
            for query, expected_days, expected_last in cases:
                with self.subTest(path=path, query=query), patch.object(main, "_fetch", return_value=[]) as fetch:
                    status, body = asyncio.run(request_export(path, query))
                    self.assertEqual(status, 200)
                    self.assertEqual(fetch.call_count, expected_days)
                    self.assertIn(expected_last, [call.args[0] for call in fetch.call_args_list])
                    if path == "/export":
                        self.assertTrue(load_workbook(io.BytesIO(body)).active["A1"].value)
                    else:
                        self.assertIn("예약 내역 없음", body.decode("utf-8-sig"))


if __name__ == "__main__":
    unittest.main()
