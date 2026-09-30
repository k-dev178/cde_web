from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
import requests as http_client
from datetime import date, datetime, timedelta
from calendar import monthrange
from pathlib import Path
from urllib.parse import quote
import io

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FRONTEND_DIR = PROJECT_ROOT / "frontend"

app = FastAPI()
app.mount("/static", StaticFiles(directory=str(FRONTEND_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(FRONTEND_DIR / "templates"))

checkin_set: set[str] = set()

BASE_URL = "https://cde.jj.ac.kr/_custom/jj/_common/app/room-reservation/logic/ajax.jsp"
WEEKDAY_KO = ["월", "화", "수", "목", "금", "토", "일"]

STUDIO_SUMMARY_LABELS = [
    "셀프(1인)스튜디오 1, 2호실(학생회관 238호)",
    "셀프(1인)스튜디오 3, 4호실(교수연구동 9층)",
    "MOOC 스튜디오",
    "하이브리드 스튜디오",
    "화상회의실(HATCH 스튜디오)",
    "HATCH 스튜디오(블랙 스튜디오)",
]


def _studio_summary(day_data: dict[date, list]) -> tuple[int, dict[str, int]]:
    """선택 기간의 예약 건수와 스튜디오 유형별 건수를 반환한다."""
    counts = {label: 0 for label in STUDIO_SUMMARY_LABELS}
    total = 0

    for items in day_data.values():
        for item in items:
            total += 1
            name = "".join(str(item.get("rpName", "")).lower().split())

            if "셀프" in name and ("교수연구동" in name or "3,4호실" in name):
                counts[STUDIO_SUMMARY_LABELS[1]] += 1
            elif "셀프" in name and ("학생회관" in name or "1,2호실" in name):
                counts[STUDIO_SUMMARY_LABELS[0]] += 1
            elif "mooc" in name:
                counts[STUDIO_SUMMARY_LABELS[2]] += 1
            elif "하이브리드" in name:
                counts[STUDIO_SUMMARY_LABELS[3]] += 1
            elif "화상회의실" in name:
                counts[STUDIO_SUMMARY_LABELS[4]] += 1
            elif "블랙" in name or "hatch스튜디오" in name:
                counts[STUDIO_SUMMARY_LABELS[5]] += 1

    return total, counts


def _fetch(d: str) -> list | None:
    try:
        r = http_client.get(BASE_URL, params={"mode": "day-list", "date": d}, timeout=10)
        return r.json().get("items", [])
    except Exception:
        return None


def _build_context() -> dict:
    today = date.today()
    d = today.strftime("%Y-%m-%d")
    items = _fetch(d)
    sorted_items = sorted(items or [], key=lambda x: x.get("rrStartTime", ""))
    return {
        "date_label": f"{today.year}년 {today.month}월 {today.day}일 ({WEEKDAY_KO[today.weekday()]})",
        "today_ym": today.strftime("%Y-%m"),
        "items": sorted_items,
        "checkin_list": list(checkin_set),
        "checkin_set": checkin_set,
        "error": items is None,
        "updated_at": datetime.now().strftime("%H:%M"),
        "total": len(sorted_items),
        "confirmed": sum(1 for i in sorted_items if i.get("rrState") == "예약완료"),
    }


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse(request=request, name="index.html", context=_build_context())


@app.get("/health", include_in_schema=False)
def health():
    return {"app": "cde-studio", "status": "ok"}


@app.get("/table", response_class=HTMLResponse)
async def table_partial(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="components/reservations/_table.html",
        context=_build_context(),
    )


@app.post("/checkin/{rr_seq}", response_class=HTMLResponse)
async def toggle_checkin(request: Request, rr_seq: str):
    if rr_seq in checkin_set:
        checkin_set.discard(rr_seq)
    else:
        checkin_set.add(rr_seq)
    return templates.TemplateResponse(
        request=request,
        name="components/reservations/_table.html",
        context=_build_context(),
    )


def _resolve_export_period(
    year: int | None,
    month: int | None,
    start_date: date | None,
    end_date: date | None,
) -> tuple[date, date, bool]:
    """월 전체 또는 양 끝 날짜를 포함하는 지정 기간을 검증한다."""
    if start_date is not None or end_date is not None:
        if start_date is None or end_date is None:
            raise HTTPException(status_code=400, detail="시작일과 종료일을 모두 지정해 주세요.")
        if start_date > end_date:
            raise HTTPException(status_code=400, detail="종료일은 시작일과 같거나 이후여야 합니다.")
        return start_date, end_date, True

    today = date.today()
    y = today.year if year is None else year
    m = today.month if month is None else month
    try:
        start = date(y, m, 1)
        end = date(y, m, monthrange(y, m)[1])
    except ValueError:
        raise HTTPException(status_code=400, detail="올바른 연도와 월을 지정해 주세요.")
    return start, end, False


def _export_period_labels(start: date, end: date, is_range: bool) -> tuple[str, str]:
    if is_range:
        return (
            f"{start:%Y.%m.%d} ~ {end:%Y.%m.%d}",
            f"{start.isoformat()}_{end.isoformat()}",
        )
    return f"{start.year}년 {start.month}월", f"{start.year}년{start.month:02d}월"


def _fetch_period(start: date, end: date) -> dict[date, list]:
    """지정 기간을 병렬 조회해 예약완료 내역을 날짜별로 반환한다."""
    from concurrent.futures import ThreadPoolExecutor, as_completed

    days = [start + timedelta(days=offset) for offset in range((end - start).days + 1)]
    day_data: dict[date, list] = {}
    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(_fetch, d.isoformat()): d for d in days}
        for future in as_completed(futures):
            day = futures[future]
            items = future.result()
            if items:
                confirmed_items = [item for item in items if item.get("rrState") == "예약완료"]
                if confirmed_items:
                    day_data[day] = sorted(
                        confirmed_items,
                        key=lambda x: x.get("rrStartTime", ""),
                    )
    return day_data


@app.get("/export")
def export_excel(
    year: int | None = None,
    month: int | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
):
    start, end, is_range = _resolve_export_period(year, month, start_date, end_date)
    period_label, filename_label = _export_period_labels(start, end, is_range)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        from fastapi.responses import JSONResponse
        return JSONResponse({"error": "openpyxl not installed"}, status_code=500)

    day_data = _fetch_period(start, end)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{start:%Y.%m.%d}~{end:%Y.%m.%d}" if is_range else period_label

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 32

    DATE_FILL  = PatternFill("solid", fgColor="1E3A8A")
    DATE_FONT  = Font(bold=True, color="FFFFFF", size=11)
    COL_FILL   = PatternFill("solid", fgColor="DBEAFE")
    COL_FONT   = Font(bold=True, color="1E40AF", size=10)
    THIN       = Side(style="thin", color="CBD5E1")
    BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    cur = 1

    for d in sorted(day_data):
        lbl = f"{d.year}년 {d.month}월 {d.day}일 ({WEEKDAY_KO[d.weekday()]})"

        # 날짜 헤더
        ws.merge_cells(f"A{cur}:C{cur}")
        cell = ws.cell(row=cur, column=1, value=lbl)
        cell.font      = DATE_FONT
        cell.fill      = DATE_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[cur].height = 20
        cur += 1

        # 컬럼 헤더
        for col, header in enumerate(["스튜디오", "이름", "사유"], 1):
            cell = ws.cell(row=cur, column=col, value=header)
            cell.font      = COL_FONT
            cell.fill      = COL_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = BORDER
        cur += 1

        # 데이터 행
        for item in day_data[d]:
            aligns = ["left", "center", "left"]
            for col, (val, align) in enumerate(zip(
                [item.get("rpName", ""), item.get("rrBooker", ""), item.get("rrPurpose", "")],
                aligns,
            ), 1):
                cell = ws.cell(row=cur, column=col, value=val)
                cell.alignment = Alignment(horizontal=align, vertical="center")
                cell.border    = BORDER
            cur += 1

        cur += 1  # 날짜 사이 빈 행

    if cur == 1:
        ws.cell(row=1, column=1, value=f"{period_label} 예약 내역 없음")
        cur = 3

    # 선택 기간의 스튜디오·기자재 대여 요약
    total, summary_counts = _studio_summary(day_data)
    ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=3)
    summary_label = period_label if is_range else f"{start.month}월"
    cell = ws.cell(row=cur, column=1, value=f"• {summary_label} 스튜디오·기자재 대여 요약(총 {total}건)")
    cell.font = Font(bold=True, size=11)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cur += 1

    for label in STUDIO_SUMMARY_LABELS:
        ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=3)
        cell = ws.cell(row=cur, column=1, value=f" - {label}: {summary_counts[label]}건")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cur += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"CDE스튜디오_{filename_label}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@app.get("/export/txt")
def export_txt(
    year: int | None = None,
    month: int | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
):
    start, end, is_range = _resolve_export_period(year, month, start_date, end_date)
    period_label, filename_label = _export_period_labels(start, end, is_range)
    day_data = _fetch_period(start, end)

    lines: list[str] = [f"■ {period_label} CDE 스튜디오 예약 현황\n"]

    for d in sorted(day_data):
        lbl = f"{d.year}년 {d.month}월 {d.day}일 ({WEEKDAY_KO[d.weekday()]})"
        lines.append(f"{'─' * 40}")
        lines.append(lbl)
        lines.append(f"{'─' * 40}")

        for item in day_data[d]:
            studio  = item.get("rpName",    "")
            booker  = item.get("rrBooker",  "")
            purpose = item.get("rrPurpose", "")
            lines.append(f"{studio} | {booker} | {purpose}")

        lines.append("")

    if not day_data:
        lines.append("예약 내역 없음")

    content = "\n".join(lines)
    buf = io.BytesIO(content.encode("utf-8-sig"))  # BOM 포함 → 메모장 한글 깨짐 방지

    filename = f"CDE스튜디오_{filename_label}.txt"
    return StreamingResponse(
        buf,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
